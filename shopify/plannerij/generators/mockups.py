"""Product mockup images for the Plannerij Shopify store (1600x1600 PNG)."""
import os
import sys

import pymupdf
from PIL import Image, ImageDraw, ImageFilter, ImageFont

BASE = "/tmp/claude-0/-home-user-frisoj/b2406013-c3ca-5da6-8b93-0de4c5b2366a/scratchpad"
FONTS = f"{BASE}/build/fonts"
OUT = f"{BASE}/img"
os.makedirs(OUT, exist_ok=True)

S = 1600
SAND = (244, 239, 230)
SAGE = (95, 122, 97)
SAGE_DARK = (70, 92, 72)
LIGHT = (232, 239, 230)
DARK = (31, 42, 36)
GREY = (138, 143, 139)
WHITE = (255, 255, 255)


def font(name, size):
    return ImageFont.truetype(f"{FONTS}/{name}.ttf", size)


def render_page(pdf, page, width):
    doc = pymupdf.open(pdf)
    p = doc[page]
    zoom = width / p.rect.width
    pix = p.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def shadow_paste(canvas, img, xy, angle=0, radius=18, blur=28, offset=(0, 26), alpha=90):
    """Paste img with rounded corners, rotation and a soft drop shadow."""
    w, h = img.size
    mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(mask).rounded_rectangle((0, 0, w - 1, h - 1), radius=radius, fill=255)
    card = Image.new("RGBA", (w, h))
    card.paste(img, (0, 0), mask)
    if angle:
        card = card.rotate(angle, expand=True, resample=Image.BICUBIC)
    sh = Image.new("RGBA", (card.width + blur * 4, card.height + blur * 4), (0, 0, 0, 0))
    sh_mask = card.split()[3]
    sh_layer = Image.new("RGBA", card.size, (20, 30, 20, alpha))
    sh.paste(sh_layer, (blur * 2, blur * 2), sh_mask)
    sh = sh.filter(ImageFilter.GaussianBlur(blur))
    canvas.alpha_composite(sh, (xy[0] - blur * 2 + offset[0], xy[1] - blur * 2 + offset[1]))
    canvas.alpha_composite(card, xy)


def base(bg=SAND, blob=True):
    im = Image.new("RGBA", (S, S), bg + (255,))
    if blob:
        d = ImageDraw.Draw(im)
        d.ellipse((S * 0.55, -S * 0.25, S * 1.35, S * 0.55), fill=LIGHT + (255,))
        d.ellipse((-S * 0.3, S * 0.7, S * 0.35, S * 1.3), fill=LIGHT + (255,))
    return im


def label(im, title, sub=None, badge=None, pos=(90, 90), dark=True):
    d = ImageDraw.Draw(im)
    # brand pill
    d.rounded_rectangle((pos[0], pos[1], pos[0] + 220, pos[1] + 54), radius=27, fill=SAGE)
    d.text((pos[0] + 110, pos[1] + 27), "Plannerij", font=font("PlayfairDisplay-600", 30), fill=WHITE, anchor="mm")
    y = pos[1] + 84
    d.text((pos[0], y), title, font=font("PlayfairDisplay-700", 76), fill=DARK if dark else WHITE)
    y += 96
    if sub:
        d.text((pos[0], y), sub, font=font("Inter-500", 34), fill=GREY if dark else LIGHT)
        y += 54
    if badge:
        f = font("Inter-600", 28)
        tw = d.textlength(badge, font=f)
        d.rounded_rectangle((pos[0], y + 10, pos[0] + tw + 44, y + 64), radius=27, outline=SAGE, width=3)
        d.text((pos[0] + 22 + tw / 2, y + 37), badge, font=f, fill=SAGE_DARK, anchor="mm")


def tablet_frame(page_img, bezel=44, radius=70):
    w, h = page_img.size
    fr = Image.new("RGBA", (w + bezel * 2, h + bezel * 2), (0, 0, 0, 0))
    d = ImageDraw.Draw(fr)
    d.rounded_rectangle((0, 0, fr.width - 1, fr.height - 1), radius=radius, fill=(28, 30, 30, 255))
    d.rounded_rectangle((6, 6, fr.width - 7, fr.height - 7), radius=radius - 4, outline=(70, 72, 72, 255), width=3)
    inner_mask = Image.new("L", (w, h), 0)
    ImageDraw.Draw(inner_mask).rounded_rectangle((0, 0, w - 1, h - 1), radius=22, fill=255)
    fr.paste(page_img, (bezel, bezel), inner_mask)
    return fr


def save(im, name):
    path = f"{OUT}/{name}.png"
    im.convert("RGB").save(path, optimize=True)
    print(path)
    return path


# ------------------------------------------------------------------ planner
PL = f"{BASE}/out/planner/Plannerij-Digitale-Planner-2027.pdf"
PA4 = f"{BASE}/out/planner/Plannerij-Printbare-Planner-2027-A4.pdf"


def planner_hero():
    im = base()
    label(im, "Digitale Planner 2027", "GoodNotes · Notability · Samsung Notes", "Direct downloaden")
    p_week = render_page(PL, 6, 1150)
    p_month = render_page(PL, 4, 1150)
    p_year = render_page(PL, 2, 1150)
    shadow_paste(im, p_year, (560, 620), angle=6)
    shadow_paste(im, p_month, (300, 720), angle=-4)
    shadow_paste(im, tablet_frame(p_week), (110, 470))
    save(im, "planner-1-hero")


def planner_pages():
    for i, (page, name, cap) in enumerate([(4, "planner-2-maand", "Maandpagina met weeknummers & feestdagen"),
                                            (6, "planner-3-week", "Weekpagina maandag t/m zondag"),
                                            (5, "planner-4-gewoontes", "Gewoontetracker per maand"),
                                            (2, "planner-5-jaar", "Jaaroverzicht met klikbare maanden")]):
        im = base(blob=False)
        d = ImageDraw.Draw(im)
        d.rectangle((0, 0, S, 150), fill=SAGE)
        d.text((80, 75), cap, font=font("Inter-600", 40), fill=WHITE, anchor="lm")
        pg = render_page(PL, page, 1440)
        shadow_paste(im, pg, (80, 260), radius=22)
        save(im, name)


def planner_printable():
    im = base()
    label(im, "Ook printbaar", "Bonus: A4 en A5 versie, 82 pagina's", "Inclusief")
    p1 = render_page(PA4, 3, 620)
    p2 = render_page(PA4, 5, 620)
    p3 = render_page(PA4, 1, 620)
    shadow_paste(im, p3, (860, 480), angle=7)
    shadow_paste(im, p2, (520, 560), angle=-5)
    shadow_paste(im, p1, (180, 640), angle=2)
    save(im, "planner-6-printbaar")


# ------------------------------------------------------------------ printables
PB = f"{BASE}/out/printables/Plannerij-Printable-Bundel-A4.pdf"


def printables_hero():
    im = base()
    label(im, "Printable Bundel", "12 ongedateerde planners & trackers", "A4 + A5 · PDF")
    pages = [(1, 900, 620, 8), (3, 560, 700, -6), (4, 200, 780, 3)]
    for pg, x, y, ang in pages:
        shadow_paste(im, render_page(PB, pg, 640), (x, y), angle=ang)
    save(im, "printables-1-hero")


def printables_grid():
    im = base(blob=False)
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, S, 150), fill=SAGE)
    d.text((80, 75), "Wat zit erin: 12 pagina's", font=font("Inter-600", 40), fill=WHITE, anchor="lm")
    pages = [1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12]
    cw, ch = 300, 424
    for i, pg in enumerate(pages):
        col, row = i % 4, i // 4
        x = 120 + col * (cw + 50)
        y = 200 + row * (ch + 36)
        shadow_paste(im, render_page(PB, pg, cw), (x, y), radius=10, blur=14, offset=(0, 10), alpha=70)
    save(im, "printables-2-overzicht")


def printables_pages():
    for pg, name, cap in [(3, "printables-3-dagplanner", "Dagplanner met tijdblokken"),
                          (4, "printables-4-gewoontetracker", "Gewoontetracker voor 31 dagen"),
                          (7, "printables-5-maaltijdplanner", "Maaltijdplanner + boodschappenlijst")]:
        im = base(blob=False)
        d = ImageDraw.Draw(im)
        d.rectangle((0, 0, S, 150), fill=SAGE)
        d.text((80, 75), cap, font=font("Inter-600", 40), fill=WHITE, anchor="lm")
        p = render_page(PB, pg, 940)
        shadow_paste(im, p, (330, 200), radius=18)
        save(im, name)


# ------------------------------------------------------------------ prompts
PR = f"{BASE}/out/prompts/Plannerij-AI-Prompt-Pack.pdf"


def prompts_hero():
    im = base()
    label(im, "AI Prompt Pack", "200 Nederlandse prompts · ChatGPT, Claude & Gemini", "PDF + CSV + TXT")
    doc = pymupdf.open(PR)
    n = doc.page_count
    pages = [(0, 900, 600, 8), (min(5, n - 1), 560, 690, -6), (min(12, n - 1), 200, 780, 3)]
    for pg, x, y, ang in pages:
        shadow_paste(im, render_page(PR, pg, 640), (x, y), angle=ang)
    save(im, "prompts-1-hero")


def prompts_pages(page_indices):
    for i, (pg, cap) in enumerate(page_indices):
        im = base(blob=False)
        d = ImageDraw.Draw(im)
        d.rectangle((0, 0, S, 150), fill=SAGE)
        d.text((80, 75), cap, font=font("Inter-600", 40), fill=WHITE, anchor="lm")
        p = render_page(PR, pg, 940)
        shadow_paste(im, p, (330, 200), radius=18)
        save(im, f"prompts-{i + 2}-pagina")


def prompts_categories():
    im = base()
    d = ImageDraw.Draw(im)
    label(im, "10 categorieën", "20 prompts per categorie, klaar om te kopiëren")
    cats = ["Ondernemen & strategie", "Marketing & social media", "Content & copywriting",
            "E-mail & klantcommunicatie", "Productiviteit & planning", "Studie & leren",
            "Carrière & solliciteren", "Geld & financiën", "Persoonlijke groei & gezondheid", "Huishouden & gezin"]
    y = 460
    for i, ct in enumerate(cats):
        col, row = i % 2, i // 2
        x = 90 + col * 720
        yy = y + row * 200
        d.rounded_rectangle((x, yy, x + 680, yy + 160), radius=28, fill=WHITE)
        d.ellipse((x + 34, yy + 50, x + 94, yy + 110), fill=LIGHT)
        d.text((x + 64, yy + 80), str(i + 1), font=font("Inter-700", 30), fill=SAGE_DARK, anchor="mm")
        d.text((x + 124, yy + 80), ct, font=font("Inter-600", 34), fill=DARK, anchor="lm")
    save(im, "prompts-6-categorieen")


# ------------------------------------------------------------------ budget (rendered from xlsx)
def sheet_image(xlsx, sheet, width=1400, max_rows=34, max_cols=10):
    """Rough visual rendering of a worksheet (fills, bold, values)."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx)
    ws = wb[sheet]
    merged = {}
    for rng in ws.merged_cells.ranges:
        merged[(rng.min_row, rng.min_col)] = rng.max_col - rng.min_col + 1
        for r_ in range(rng.min_row, rng.max_row + 1):
            for c_ in range(rng.min_col, rng.max_col + 1):
                if (r_, c_) != (rng.min_row, rng.min_col):
                    merged[(r_, c_)] = 0
    colw = []
    for ci in range(1, max_cols + 1):
        letter = openpyxl.utils.get_column_letter(ci)
        w = ws.column_dimensions[letter].width or 10
        colw.append(max(4, min(w, 40)) * 7.2)
    total_w = sum(colw)
    scale = width / total_w
    rowh = 26
    H = int(rowh * (max_rows + 1) * scale) + 40
    im = Image.new("RGB", (width, H), WHITE)
    d = ImageDraw.Draw(im)
    f = ImageFont.truetype(f"{FONTS}/Inter-400.ttf", int(13 * scale * 1.35))
    fb = ImageFont.truetype(f"{FONTS}/Inter-600.ttf", int(13 * scale * 1.35))
    y = 0
    for ri in range(1, max_rows + 1):
        x = 0
        rh = rowh * scale
        for ci in range(1, max_cols + 1):
            cell = ws.cell(row=ri, column=ci)
            cw = colw[ci - 1] * scale
            span = merged.get((ri, ci), 1)
            if span == 0:
                x += cw
                continue
            if span > 1:
                cw = sum(colw[ci - 1:ci - 1 + span]) * scale
            fill = None
            try:
                rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
                if isinstance(rgb, str) and len(rgb) >= 6 and rgb != "00000000":
                    fill = tuple(int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4))
            except Exception:
                fill = None
            if fill:
                d.rectangle((x, y, x + cw, y + rh), fill=fill)
            d.rectangle((x, y, x + cw, y + rh), outline=(235, 235, 235))
            v = cell.value
            if v is not None and not (isinstance(v, str) and v.startswith("=")):
                s = str(v)
                if isinstance(v, (int, float)):
                    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
                col = DARK
                try:
                    fc = cell.font.color.rgb if cell.font and cell.font.color else None
                    if isinstance(fc, str) and len(fc) >= 6 and fc not in ("00000000",):
                        col = tuple(int(fc[-6:][i:i + 2], 16) for i in (0, 2, 4))
                except Exception:
                    pass
                fnt = fb if (cell.font and cell.font.bold) else f
                maxch = int(cw / (7.2 * scale)) + 1
                if len(s) > maxch:
                    s = s[:max(1, maxch - 1)] + "…"
                d.text((x + 6 * scale, y + rh / 2), s, font=fnt, fill=col, anchor="lm")
            elif isinstance(v, str) and v.startswith("="):
                nf = cell.number_format or ""
                shown = "€ 0,00" if "€" in nf else ("0%" if "%" in nf else ("" if "Instellingen" in v else "0"))
                if "Instellingen" in v:
                    shown = ""
                d.text((x + 6 * scale, y + rh / 2), shown, font=f, fill=GREY, anchor="lm")
            x += cw
        y += rh
    # sheet tab strip
    return im


def budget_images(xlsx, sheets):
    im = base()
    label(im, "Budgetplanner", "Excel · Google Sheets · Numbers", "Nederlands · €")
    imgs = [sheet_image(xlsx, s, 1200, 30, 9) for s in sheets[:2]]
    shadow_paste(im, imgs[1], (520, 620), angle=5, radius=14)
    shadow_paste(im, imgs[0], (120, 560), angle=-3, radius=14)
    save(im, "budget-1-hero")
    for i, s in enumerate(sheets):
        im = base(blob=False)
        d = ImageDraw.Draw(im)
        d.rectangle((0, 0, S, 150), fill=SAGE)
        d.text((80, 75), f"Tabblad: {s}", font=font("Inter-600", 40), fill=WHITE, anchor="lm")
        shadow_paste(im, sheet_image(xlsx, s, 1440, 36, 10), (80, 230), radius=14)
        save(im, f"budget-{i + 2}-{s.lower().replace(' ', '-')}")


# ------------------------------------------------------------------ bundle
def bundle_hero():
    im = base()
    label(im, "Complete Bundel 2027", "Planner + Budgetplanner + Prompt Pack + Printables", "Bespaar ruim 30%")
    items = [(render_page(PL, 4, 700), (60, 560), -5), (render_page(PB, 3, 380), (1140, 520), 6),
             (render_page(PR, 0, 380), (860, 900), -7)]
    for img, xy, ang in items:
        shadow_paste(im, img, xy, angle=ang)
    if os.path.exists(f"{BASE}/out/budget/Plannerij-Budgetplanner.xlsx"):
        shadow_paste(im, sheet_image(f"{BASE}/out/budget/Plannerij-Budgetplanner.xlsx", "Jaaroverzicht", 760, 24, 8), (400, 1000), angle=3, radius=14)
    save(im, "bundel-1-hero")


def collection_banner():
    im = Image.new("RGBA", (2000, 1000), SAGE + (255,))
    d = ImageDraw.Draw(im)
    d.ellipse((1300, -300, 2300, 700), fill=SAGE_DARK + (255,))
    d.text((120, 300), "Plannerij", font=font("PlayfairDisplay-700", 150), fill=WHITE)
    d.text((124, 500), "Rust in je hoofd begint met overzicht.", font=font("Inter-500", 54), fill=LIGHT)
    d.text((124, 590), "Digitale planners, budgetplanners en AI-prompts.", font=font("Inter-400", 36), fill=LIGHT)
    d.text((124, 645), "Nederlands · Direct te downloaden", font=font("Inter-400", 36), fill=LIGHT)
    pg = render_page(PL, 6, 820)
    shadow_paste(im, tablet_frame(pg, bezel=30, radius=50), (1180, 220), angle=-4)
    save(im, "hero-banner")
    # square logo
    lg = Image.new("RGBA", (800, 800), SAGE + (255,))
    d = ImageDraw.Draw(lg)
    d.text((400, 380), "P", font=font("PlayfairDisplay-700", 420), fill=WHITE, anchor="mm")
    d.text((400, 640), "PLANNERIJ", font=font("Inter-600", 64), fill=LIGHT, anchor="mm")
    save(lg, "logo-square")
    # wordmark logo (transparent-ish on white)
    wm = Image.new("RGBA", (1200, 300), (0, 0, 0, 0))
    d = ImageDraw.Draw(wm)
    d.text((0, 150), "Plannerij", font=font("PlayfairDisplay-700", 190), fill=SAGE_DARK, anchor="lm")
    wm.save(f"{OUT}/logo-wordmark.png")
    print(f"{OUT}/logo-wordmark.png")


if __name__ == "__main__":
    what = sys.argv[1:] or ["planner", "printables", "brand"]
    if "planner" in what:
        planner_hero(); planner_pages(); planner_printable()
    if "printables" in what:
        printables_hero(); printables_grid(); printables_pages()
    if "prompts" in what:
        prompts_hero(); prompts_categories()
        doc = pymupdf.open(PR)
        n = doc.page_count
        prompts_pages([(min(4, n - 1), "Overzichtelijke promptkaarten"), (min(9, n - 1), "Elke prompt met een tip"),
                       (1, "Zo gebruik je het pakket")])
    if "budget" in what:
        budget_images(f"{BASE}/out/budget/Plannerij-Budgetplanner.xlsx", ["Jaaroverzicht", "Januari", "Spaardoelen", "Abonnementen"])
    if "bundle" in what:
        bundle_hero()
    if "brand" in what:
        collection_banner()
