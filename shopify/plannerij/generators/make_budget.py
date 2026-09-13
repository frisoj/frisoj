#!/usr/bin/env python3
"""Genereert de Plannerij Budgetplanner (xlsx) met openpyxl."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path("/tmp/claude-0/-home-user-frisoj/b2406013-c3ca-5da6-8b93-0de4c5b2366a/scratchpad/out/budget")
OUT_FILE = OUT_DIR / "Plannerij-Budgetplanner.xlsx"

# ---------------------------------------------------------------- huisstijl
GREEN = "5F7A61"
LIGHT = "E8EFE6"
SAND = "F4EFE6"
DARK = "1F2A24"
GREY = "D9D9D9"
RED = "B5482B"
WHITE = "FFFFFF"

FONT = "Calibri"
EUR = '"€ "#,##0.00'
PCT = "0%"
DATE = "dd-mm-yyyy"

f_base = Font(name=FONT, size=11, color=DARK)
f_bold = Font(name=FONT, size=11, color=DARK, bold=True)
f_title = Font(name=FONT, size=20, color=GREEN, bold=True)
f_logo = Font(name=FONT, size=14, color=GREEN, bold=True)
f_sub = Font(name=FONT, size=11, color="6B7570", italic=True)
f_head = Font(name=FONT, size=11, color=WHITE, bold=True)
f_small = Font(name=FONT, size=10, color="6B7570")
f_kpi = Font(name=FONT, size=16, color=GREEN, bold=True)
f_red = Font(name=FONT, size=11, color=RED, bold=True)

fill_green = PatternFill("solid", fgColor=GREEN)
fill_light = PatternFill("solid", fgColor=LIGHT)
fill_sand = PatternFill("solid", fgColor=SAND)
fill_white = PatternFill("solid", fgColor=WHITE)

thin = Side(style="thin", color=GREY)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
border_bottom = Border(bottom=Side(style="medium", color=GREEN))

al_left = Alignment(horizontal="left", vertical="center")
al_center = Alignment(horizontal="center", vertical="center")
al_right = Alignment(horizontal="right", vertical="center")
al_wrap = Alignment(horizontal="left", vertical="top", wrap_text=True)

MONTHS = ["Januari", "Februari", "Maart", "April", "Mei", "Juni",
          "Juli", "Augustus", "September", "Oktober", "November", "December"]

INKOMEN = ["Salaris", "Toeslagen", "Freelance", "Overig"]
VAST = ["Huur/Hypotheek", "Energie", "Water", "Zorgverzekering", "Overige verzekeringen",
        "Internet & TV", "Telefoon", "Abonnementen", "Gemeentelijke belastingen",
        "Auto/OV", "Kinderopvang", "Aflossing schulden"]
VARIABEL = ["Boodschappen", "Uit eten", "Kleding", "Persoonlijke verzorging",
            "Vervoer/Brandstof", "Cadeaus", "Vrije tijd", "Vakantie", "Huis & Tuin",
            "Gezondheid", "Overig"]
SPAAR = ["Noodbuffer", "Vakantie", "Grote aankoop", "Pensioen/beleggen"]

# Posities in Instellingen (rij 9 = eerste item)
SET_ROW0 = 9
SET_COL = {"ink": "B", "vast": "D", "var": "F", "spaar": "H"}


# ---------------------------------------------------------------- helpers
def prep(ws, widths, freeze="A4"):
    ws.sheet_view.showGridLines = False
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    if freeze:
        ws.freeze_panes = freeze
    ws.sheet_properties.tabColor = GREEN


def cell(ws, ref, value=None, font=f_base, fill=None, fmt=None, align=None, brd=None):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if brd:
        c.border = brd
    return c


def title(ws, text, sub=None):
    cell(ws, "B1", "Plannerij", f_logo)
    cell(ws, "B2", text, f_title)
    if sub:
        cell(ws, "B3", sub, f_sub)
    ws.row_dimensions[2].height = 30


def header_row(ws, row, col_start, labels, fill=fill_green, font=f_head):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col_start + i, value=lab)
        c.font = font
        c.fill = fill
        c.border = border
        c.alignment = al_center if i else al_left


def input_cell(ws, ref, fmt=None, align=None):
    return cell(ws, ref, None, f_base, fill_light, fmt, align or al_right, border)


def formula_cell(ws, ref, formula, fmt=None, bold=False, align=None, fill=fill_white):
    return cell(ws, ref, formula, f_bold if bold else f_base, fill, fmt, align or al_right, border)


def label_cell(ws, ref, text, bold=False, fill=None):
    return cell(ws, ref, text, f_bold if bold else f_base, fill, None, al_left, border)


def red_if_negative(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], font=f_red))


# ---------------------------------------------------------------- Start
def build_start(wb):
    ws = wb.active
    ws.title = "Start"
    prep(ws, {"A": 3, "B": 6, "C": 95}, freeze=None)
    cell(ws, "B1", "Plannerij", f_logo)
    cell(ws, "B2", "Budgetplanner", f_title)
    cell(ws, "B3", "Grip op je geld – maand voor maand. Werkt in Excel, Google Sheets en Numbers.", f_sub)
    ws.row_dimensions[2].height = 32

    r = 5
    cell(ws, f"B{r}", "Zo gebruik je dit bestand", f_bold)
    ws[f"B{r}"].font = Font(name=FONT, size=13, color=GREEN, bold=True)
    steps = [
        "Ga naar het tabblad Instellingen. Vul de naam van je huishouden en je startmaand in en pas de categorienamen aan naar jouw situatie. Alle maandtabbladen nemen die namen automatisch over.",
        "Open het tabblad van de huidige maand (Januari t/m December) en vul in de kolom Begroot in wat je verwacht aan inkomen, vaste lasten, variabele uitgaven en sparen.",
        "Vul in de loop van de maand je werkelijke inkomen en vaste lasten in (groene cellen in de kolom Werkelijk).",
        "Noteer variabele uitgaven (boodschappen, uit eten, kleding, …) in het Uitgavenlogboek rechts op het maandblad: datum, omschrijving, categorie (keuzelijst) en bedrag. De kolom Werkelijk bij Variabele uitgaven telt dit automatisch op per categorie (SUMIF).",
        "Vul onder Sparen in wat je werkelijk naar elk spaardoel hebt overgemaakt. Het tabblad Spaardoelen houdt de voortgang bij.",
        "Bekijk het samenvattingskader rechtsboven op het maandblad: Totaal inkomen, Totaal uitgaven, Gespaard en Over/tekort. Rood betekent dat je meer hebt uitgegeven dan er binnenkwam.",
        "Kijk regelmatig op het Jaaroverzicht voor het totaalbeeld: inkomen versus uitgaven per maand, je spaarquote en de grafiek.",
        "Gebruik de extra tabbladen Schulden, Abonnementen en Spaarpotjes 52 weken om schulden af te bouwen, abonnementen op te ruimen en spelenderwijs te sparen.",
    ]
    for i, s in enumerate(steps, 1):
        r += 1
        cell(ws, f"B{r}", f"{i}.", f_bold, align=Alignment(horizontal="right", vertical="top"))
        cell(ws, f"C{r}", s, f_base, align=al_wrap)
        ws.row_dimensions[r].height = 34

    r += 2
    cell(ws, f"B{r}", "Legenda", Font(name=FONT, size=13, color=GREEN, bold=True))
    r += 1
    cell(ws, f"B{r}", None, f_base, fill_light, brd=border)
    cell(ws, f"C{r}", "Groene cel = invulcel. Hier vul je je eigen cijfers in.", f_base, align=al_left)
    r += 1
    cell(ws, f"B{r}", None, f_base, fill_white, brd=border)
    cell(ws, f"C{r}", "Witte cel = formule. Niet aanpassen; deze rekent automatisch.", f_base, align=al_left)
    r += 1
    cell(ws, f"B{r}", None, f_base, fill_sand, brd=border)
    cell(ws, f"C{r}", "Zandkleurige cel = subtotaal of samenvatting.", f_base, align=al_left)
    r += 1
    cell(ws, f"C{r}", "Voorbeeld-bedragen? Vul je eigen cijfers in in de groene cellen. Het bestand is bewust leeg opgeleverd, zodat je met een schone lei begint.", f_bold, align=al_wrap)
    ws.row_dimensions[r].height = 30

    r += 2
    cell(ws, f"B{r}", "Tips", Font(name=FONT, size=13, color=GREEN, bold=True))
    tips = [
        "Verschil-kolom: bij Inkomen is dat Werkelijk – Begroot (positief = meer binnen dan verwacht). Bij uitgaven is het Begroot – Werkelijk (positief = binnen budget gebleven, rood = overschreden).",
        "Categorie hernoemen? Doe dat alléén op het tabblad Instellingen; alle maanden en keuzelijsten volgen automatisch.",
        "Spaarquote = Gespaard ÷ Inkomen. Een spaarquote van 10–20% is een mooi streven; alles boven 0% is winst.",
        "Bouw eerst een noodbuffer van 3 maanden vaste lasten op vóór je aan grotere spaardoelen begint.",
        "Vul het logboek wekelijks in (bijvoorbeeld op zondagavond) – dat kost 5 minuten en houdt je overzicht actueel.",
        "Wil je een tabblad beschermen tegen per ongeluk overschrijven? Gebruik 'Blad beveiligen' in Excel of Google Sheets; de groene cellen blijven dan invulbaar als je ze eerst ontgrendelt.",
        "Google Sheets: upload het bestand via Bestand > Importeren. Numbers: openen werkt direct; formules blijven behouden.",
    ]
    for t in tips:
        r += 1
        cell(ws, f"B{r}", "•", f_bold, align=Alignment(horizontal="right", vertical="top"))
        cell(ws, f"C{r}", t, f_base, align=al_wrap)
        ws.row_dimensions[r].height = 32

    r += 2
    cell(ws, f"C{r}", "Werkt in Excel, Google Sheets en Numbers.  ·  © Plannerij", f_small)
    return ws


# ---------------------------------------------------------------- Instellingen
def build_instellingen(wb):
    ws = wb.create_sheet("Instellingen")
    prep(ws, {"A": 3, "B": 30, "C": 24, "D": 30, "E": 3, "F": 30, "G": 3, "H": 30, "I": 3}, freeze="A4")
    title(ws, "Instellingen", "Vul de groene cellen in. Categorienamen die je hier wijzigt, worden overal overgenomen.")

    label_cell(ws, "B5", "Naam huishouden", bold=True)
    input_cell(ws, "C5", align=al_left)
    label_cell(ws, "B6", "Startmaand (datum, bv. 01-01-2027)", bold=True)
    input_cell(ws, "C6", fmt="mmmm yyyy", align=al_left)
    label_cell(ws, "B7", "Jaar", bold=True)
    formula_cell(ws, "C7", '=IF(C6="","",YEAR(C6))', fmt="0", align=al_left)

    header_row(ws, 9 - 1, 2, ["Inkomen"])
    header_row(ws, 9 - 1, 4, ["Vaste lasten"])
    header_row(ws, 9 - 1, 6, ["Variabele uitgaven"])
    header_row(ws, 9 - 1, 8, ["Spaardoelen"])
    for col, items in (("B", INKOMEN), ("D", VAST), ("F", VARIABEL), ("H", SPAAR)):
        for i, name in enumerate(items):
            c = input_cell(ws, f"{col}{SET_ROW0 + i}", align=al_left)
            c.value = name

    r = SET_ROW0 + len(VAST) + 2
    cell(ws, f"B{r}", "Tip: hernoem categorieën gerust, maar laat de volgorde en het aantal rijen staan. "
                      "De keuzelijst in het Uitgavenlogboek gebruikt de lijst 'Variabele uitgaven'.", f_small,
         align=al_wrap)
    ws.merge_cells(f"B{r}:H{r}")
    ws.row_dimensions[r].height = 30
    return ws


# ---------------------------------------------------------------- Maandblad
# rij-layout (vast voor alle maanden)
R_INK_H = 4
R_INK0 = 5
R_INK_T = R_INK0 + len(INKOMEN)            # 9
R_VAST_H = R_INK_T + 2                      # 11
R_VAST0 = R_VAST_H + 1                      # 12
R_VAST_T = R_VAST0 + len(VAST)              # 24
R_VAR_H = R_VAST_T + 2                      # 26
R_VAR0 = R_VAR_H + 1                        # 27
R_VAR_T = R_VAR0 + len(VARIABEL)            # 38
R_SP_H = R_VAR_T + 2                        # 40
R_SP0 = R_SP_H + 1                          # 41
R_SP_T = R_SP0 + len(SPAAR)                 # 45
# samenvatting
R_SUM_H = 4
# logboek
R_LOG_TITLE = 10
R_LOG_H = 11
R_LOG0 = 12
LOG_ROWS = 40
R_LOG_END = R_LOG0 + LOG_ROWS - 1           # 51
R_LOG_T = R_LOG_END + 1                     # 52


def build_month(wb, name):
    ws = wb.create_sheet(name)
    prep(ws, {"A": 2, "B": 28, "C": 14, "D": 14, "E": 14, "F": 3,
              "G": 14, "H": 32, "I": 26, "J": 14}, freeze="A4")
    title(ws, name)
    ws["B3"] = '=IF(Instellingen!$C$5="","",Instellingen!$C$5)'
    ws["B3"].font = f_sub

    log_cat = f"$I${R_LOG0}:$I${R_LOG_END}"
    log_amt = f"$J${R_LOG0}:$J${R_LOG_END}"

    def block(r_head, r0, items, set_col, label, kind):
        header_row(ws, r_head, 2, [label, "Begroot", "Werkelijk", "Verschil"])
        for i, _ in enumerate(items):
            r = r0 + i
            formula_cell(ws, f"B{r}", f"=Instellingen!${set_col}${SET_ROW0 + i}", align=al_left)
            input_cell(ws, f"C{r}", EUR)
            if kind == "var":
                formula_cell(ws, f"D{r}", f"=SUMIF({log_cat},B{r},{log_amt})", EUR)
            else:
                input_cell(ws, f"D{r}", EUR)
            if kind in ("ink", "sp"):
                formula_cell(ws, f"E{r}", f"=D{r}-C{r}", EUR)
            else:
                formula_cell(ws, f"E{r}", f"=C{r}-D{r}", EUR)
        rt = r0 + len(items)
        label_cell(ws, f"B{rt}", "Subtotaal", bold=True, fill=fill_sand)
        for col in "CDE":
            formula_cell(ws, f"{col}{rt}", f"=SUM({col}{r0}:{col}{rt - 1})", EUR, bold=True, fill=fill_sand)
        red_if_negative(ws, f"E{r0}:E{rt}")

    block(R_INK_H, R_INK0, INKOMEN, SET_COL["ink"], "A · Inkomen", "ink")
    block(R_VAST_H, R_VAST0, VAST, SET_COL["vast"], "B · Vaste lasten", "vast")
    block(R_VAR_H, R_VAR0, VARIABEL, SET_COL["var"], "C · Variabele uitgaven", "var")
    block(R_SP_H, R_SP0, SPAAR, SET_COL["spaar"], "D · Sparen", "sp")
    cell(ws, f"B{R_SP_T + 1}", "Werkelijk bij variabele uitgaven wordt automatisch gevuld vanuit het Uitgavenlogboek.", f_small)

    # samenvatting rechtsboven (G:H = label, I = Begroot, J = Werkelijk)
    header_row(ws, R_SUM_H, 7, ["Samenvatting", "", "Begroot", "Werkelijk"])
    ws.merge_cells(start_row=R_SUM_H, start_column=7, end_row=R_SUM_H, end_column=8)
    rows = [
        ("Totaal inkomen", f"=C{R_INK_T}", f"=D{R_INK_T}"),
        ("Totaal uitgaven", f"=C{R_VAST_T}+C{R_VAR_T}", f"=D{R_VAST_T}+D{R_VAR_T}"),
        ("Gespaard", f"=C{R_SP_T}", f"=D{R_SP_T}"),
    ]
    for i, (lab, fb, fw) in enumerate(rows):
        r = R_SUM_H + 1 + i
        label_cell(ws, f"G{r}", lab, fill=fill_sand)
        cell(ws, f"H{r}", None, f_base, fill_sand, brd=border)
        ws.merge_cells(f"G{r}:H{r}")
        formula_cell(ws, f"I{r}", fb, EUR, fill=fill_sand)
        formula_cell(ws, f"J{r}", fw, EUR, fill=fill_sand)
    r = R_SUM_H + 4
    label_cell(ws, f"G{r}", "Over / tekort", bold=True, fill=fill_sand)
    cell(ws, f"H{r}", None, f_base, fill_sand, brd=border)
    ws.merge_cells(f"G{r}:H{r}")
    formula_cell(ws, f"I{r}", f"=I{r-3}-I{r-2}-I{r-1}", EUR, bold=True, fill=fill_sand)
    formula_cell(ws, f"J{r}", f"=J{r-3}-J{r-2}-J{r-1}", EUR, bold=True, fill=fill_sand)
    red_if_negative(ws, f"I{r}:J{r}")

    # logboek
    cell(ws, f"G{R_LOG_TITLE}", "Uitgavenlogboek", Font(name=FONT, size=13, color=GREEN, bold=True))
    cell(ws, f"I{R_LOG_TITLE}", "Kies een categorie uit de lijst; het bedrag telt automatisch op bij Variabele uitgaven.", f_small)
    header_row(ws, R_LOG_H, 7, ["Datum", "Omschrijving", "Categorie", "Bedrag"])
    for r in range(R_LOG0, R_LOG_END + 1):
        input_cell(ws, f"G{r}", DATE, al_center)
        input_cell(ws, f"H{r}", align=al_left)
        input_cell(ws, f"I{r}", align=al_left)
        input_cell(ws, f"J{r}", EUR)
    label_cell(ws, f"G{R_LOG_T}", "Totaal logboek", bold=True, fill=fill_sand)
    ws.merge_cells(f"G{R_LOG_T}:I{R_LOG_T}")
    formula_cell(ws, f"J{R_LOG_T}", f"=SUM(J{R_LOG0}:J{R_LOG_END})", EUR, bold=True, fill=fill_sand)

    dv = DataValidation(
        type="list",
        formula1=f"=Instellingen!${SET_COL['var']}${SET_ROW0}:${SET_COL['var']}${SET_ROW0 + len(VARIABEL) - 1}",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Onbekende categorie",
        error="Kies een categorie uit de lijst (aanpassen kan op het tabblad Instellingen).")
    ws.add_data_validation(dv)
    dv.add(f"I{R_LOG0}:I{R_LOG_END}")
    return ws


# ---------------------------------------------------------------- Jaaroverzicht
def build_jaar(wb):
    ws = wb.create_sheet("Jaaroverzicht")
    prep(ws, {"A": 3, "B": 16, "C": 16, "D": 16, "E": 18, "F": 14, "G": 16, "H": 16, "I": 14}, freeze="A9")
    title(ws, "Jaaroverzicht", "Werkelijke cijfers uit de twaalf maandtabbladen. Alles rekent automatisch.")
    ws["D3"] = '=IF(Instellingen!$C$5="","",Instellingen!$C$5)'
    ws["D3"].font = f_sub

    R_H = 8
    R0 = 9
    R_T = R0 + 12   # 21
    R_A = R_T + 1   # 22

    # KPI's
    kpis = [
        ("Totaal inkomen", f"=C{R_T}", EUR),
        ("Totaal uitgaven", f"=G{R_T}", EUR),
        ("Totaal gespaard", f"=F{R_T}", EUR),
        ("Gemiddelde spaarquote", f"=I{R_A}", PCT),
    ]
    for i, (lab, f, fmt) in enumerate(kpis):
        c1 = 2 + i * 2
        c2 = c1 + 1
        l1 = get_column_letter(c1)
        l2 = get_column_letter(c2)
        cell(ws, f"{l1}5", lab, f_small, fill_sand, align=al_center, brd=border)
        cell(ws, f"{l2}5", None, f_small, fill_sand, brd=border)
        cell(ws, f"{l1}6", f, f_kpi, fill_sand, fmt, al_center, border)
        cell(ws, f"{l2}6", None, f_kpi, fill_sand, brd=border)
        ws.merge_cells(f"{l1}5:{l2}5")
        ws.merge_cells(f"{l1}6:{l2}6")
    ws.row_dimensions[6].height = 30

    header_row(ws, R_H, 2, ["Maand", "Inkomen", "Vaste lasten", "Variabele uitgaven",
                            "Sparen", "Totaal uitgaven", "Resultaat", "Spaarquote"])
    for i, m in enumerate(MONTHS):
        r = R0 + i
        label_cell(ws, f"B{r}", m, bold=True)
        formula_cell(ws, f"C{r}", f"='{m}'!D{R_INK_T}", EUR)
        formula_cell(ws, f"D{r}", f"='{m}'!D{R_VAST_T}", EUR)
        formula_cell(ws, f"E{r}", f"='{m}'!D{R_VAR_T}", EUR)
        formula_cell(ws, f"F{r}", f"='{m}'!D{R_SP_T}", EUR)
        formula_cell(ws, f"G{r}", f"=D{r}+E{r}", EUR)
        formula_cell(ws, f"H{r}", f"=C{r}-G{r}-F{r}", EUR)
        formula_cell(ws, f"I{r}", f"=IFERROR(F{r}/C{r},0)", PCT)
    label_cell(ws, f"B{R_T}", "Totaal", bold=True, fill=fill_sand)
    label_cell(ws, f"B{R_A}", "Gemiddelde", bold=True, fill=fill_sand)
    for col in "CDEFGH":
        formula_cell(ws, f"{col}{R_T}", f"=SUM({col}{R0}:{col}{R_T - 1})", EUR, bold=True, fill=fill_sand)
        formula_cell(ws, f"{col}{R_A}", f"=AVERAGE({col}{R0}:{col}{R_T - 1})", EUR, bold=True, fill=fill_sand)
    formula_cell(ws, f"I{R_T}", f"=IFERROR(F{R_T}/C{R_T},0)", PCT, bold=True, fill=fill_sand)
    formula_cell(ws, f"I{R_A}", f"=IFERROR(F{R_T}/C{R_T},0)", PCT, bold=True, fill=fill_sand)
    red_if_negative(ws, f"H{R0}:H{R_A}")
    cell(ws, f"B{R_A + 1}", "Resultaat = Inkomen – Totaal uitgaven – Sparen. Spaarquote = Sparen ÷ Inkomen.", f_small)

    # grafiek
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = "Inkomen vs. uitgaven per maand"
    chart.y_axis.title = "€"
    chart.y_axis.numFmt = "#,##0"
    chart.x_axis.title = None
    chart.height = 9
    chart.width = 24
    chart.style = 10
    inc = Reference(ws, min_col=3, min_row=R_H, max_row=R_T - 1)
    exp = Reference(ws, min_col=7, min_row=R_H, max_row=R_T - 1)
    chart.add_data(inc, titles_from_data=True)
    chart.add_data(exp, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=R0, max_row=R_T - 1))
    chart.series[0].graphicalProperties.solidFill = GREEN
    chart.series[0].graphicalProperties.line.solidFill = GREEN
    chart.series[1].graphicalProperties.solidFill = "C9B79C"
    chart.series[1].graphicalProperties.line.solidFill = "C9B79C"
    chart.legend.position = "b"
    ws.add_chart(chart, f"B{R_A + 3}")
    return ws


# ---------------------------------------------------------------- Spaardoelen
def build_spaardoelen(wb):
    ws = wb.create_sheet("Spaardoelen")
    prep(ws, {"A": 3, "B": 26, "C": 16, "D": 18, "E": 16, "F": 16, "G": 12, "H": 16, "I": 16, "J": 18}, freeze="A6")
    title(ws, "Spaardoelen", "Doelbedrag, startbedrag en streefdatum vul je in; de rest rekent automatisch.")
    header_row(ws, 5, 2, ["Spaardoel", "Doelbedrag", "Al gespaard vóór dit jaar",
                          "Gespaard dit jaar", "Totaal gespaard", "% behaald",
                          "Nog te gaan", "Streefdatum", "Nodig per maand"])
    ws.row_dimensions[5].height = 32
    for c in range(2, 11):
        ws.cell(row=5, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for i, _ in enumerate(SPAAR):
        r = 6 + i
        formula_cell(ws, f"B{r}", f"=Instellingen!$H${SET_ROW0 + i}", align=al_left)
        input_cell(ws, f"C{r}", EUR)
        input_cell(ws, f"D{r}", EUR)
        som = "+".join(f"'{m}'!D{R_SP0 + i}" for m in MONTHS)
        formula_cell(ws, f"E{r}", f"={som}", EUR)
        formula_cell(ws, f"F{r}", f"=D{r}+E{r}", EUR)
        formula_cell(ws, f"G{r}", f"=IFERROR(MIN(1,F{r}/C{r}),0)", PCT)
        formula_cell(ws, f"H{r}", f"=MAX(0,C{r}-F{r})", EUR)
        input_cell(ws, f"I{r}", DATE, al_center)
        months = f"MAX(1,(YEAR(I{r})-YEAR(TODAY()))*12+MONTH(I{r})-MONTH(TODAY()))"
        formula_cell(ws, f"J{r}", f'=IF(I{r}="",0,IFERROR(ROUND(H{r}/{months},2),0))', EUR)
    rt = 6 + len(SPAAR)
    label_cell(ws, f"B{rt}", "Totaal", bold=True, fill=fill_sand)
    for col in "CDEFH":
        formula_cell(ws, f"{col}{rt}", f"=SUM({col}6:{col}{rt - 1})", EUR, bold=True, fill=fill_sand)
    formula_cell(ws, f"G{rt}", f"=IFERROR(MIN(1,F{rt}/C{rt}),0)", PCT, bold=True, fill=fill_sand)
    cell(ws, f"I{rt}", None, f_base, fill_sand, brd=border)
    formula_cell(ws, f"J{rt}", f"=SUM(J6:J{rt - 1})", EUR, bold=True, fill=fill_sand)
    ws.conditional_formatting.add(f"G6:G{rt - 1}",
                                  DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1,
                                              color=GREEN, showValue=True))
    cell(ws, f"B{rt + 2}", "Gespaard dit jaar = som van de kolom Werkelijk onder 'D · Sparen' op alle maandtabbladen. "
                            "Nodig per maand = Nog te gaan ÷ aantal maanden tot de streefdatum.", f_small)
    return ws


# ---------------------------------------------------------------- Schulden
def build_schulden(wb):
    ws = wb.create_sheet("Schulden")
    prep(ws, {"A": 3, "B": 24, "C": 22, "D": 18, "E": 12, "F": 16, "G": 16, "H": 16, "I": 14, "J": 18}, freeze="A6")
    title(ws, "Schulden", "Overzicht van leningen en schulden. Groene cellen invullen; restant en einddatum rekenen automatisch.")
    header_row(ws, 5, 2, ["Schuld", "Schuldeiser", "Oorspronkelijk bedrag", "Rente % / jaar",
                          "Maandbedrag", "Al afgelost", "Restant", "Maanden te gaan", "Verwachte einddatum"])
    ws.row_dimensions[5].height = 32
    for c in range(2, 11):
        ws.cell(row=5, column=c).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    n = 8
    for i in range(n):
        r = 6 + i
        input_cell(ws, f"B{r}", align=al_left)
        input_cell(ws, f"C{r}", align=al_left)
        input_cell(ws, f"D{r}", EUR)
        input_cell(ws, f"E{r}", "0.0%")
        input_cell(ws, f"F{r}", EUR)
        input_cell(ws, f"G{r}", EUR)
        formula_cell(ws, f"H{r}", f"=MAX(0,D{r}-G{r})", EUR)
        nper = f"IF(E{r}=0,H{r}/F{r},NPER(E{r}/12,-F{r},H{r}))"
        formula_cell(ws, f"I{r}",
                     f'=IF(AND(ISNUMBER(F{r}),F{r}>0,H{r}>0),IFERROR(ROUNDUP({nper},0),"n.v.t."),"")', "0", al_center)
        formula_cell(ws, f"J{r}",
                     f'=IF(ISNUMBER(I{r}),DATE(YEAR(TODAY()),MONTH(TODAY())+I{r},1),"")', "mmm yyyy", al_center)
    rt = 6 + n
    label_cell(ws, f"B{rt}", "Totaal", bold=True, fill=fill_sand)
    cell(ws, f"C{rt}", None, f_base, fill_sand, brd=border)
    for col in "DFGH":
        formula_cell(ws, f"{col}{rt}", f"=SUM({col}6:{col}{rt - 1})", EUR, bold=True, fill=fill_sand)
    for col in "EIJ":
        cell(ws, f"{col}{rt}", None, f_base, fill_sand, brd=border)
    r = rt + 2
    cell(ws, f"B{r}", "Sneeuwbal of lawine?", Font(name=FONT, size=13, color=GREEN, bold=True))
    cell(ws, f"B{r + 1}", "Sneeuwbalmethode: los eerst de kleinste schuld volledig af en gebruik dat vrijgekomen maandbedrag "
                          "voor de volgende – snelle succeservaringen houden je gemotiveerd.", f_base, align=al_wrap)
    cell(ws, f"B{r + 2}", "Lawinemethode: begin bij de schuld met de hoogste rente – wiskundig het goedkoopst, "
                          "omdat je de minste rente betaalt.", f_base, align=al_wrap)
    for rr in (r + 1, r + 2):
        ws.merge_cells(f"B{rr}:J{rr}")
        ws.row_dimensions[rr].height = 32
    cell(ws, f"B{r + 3}", "Maanden te gaan is berekend met NPER op basis van rente, maandbedrag en restant. "
                          "'n.v.t.' betekent dat het maandbedrag lager is dan de maandelijkse rente.", f_small)
    return ws


# ---------------------------------------------------------------- Abonnementen
def build_abonnementen(wb):
    ws = wb.create_sheet("Abonnementen")
    prep(ws, {"A": 3, "B": 30, "C": 16, "D": 16, "E": 12, "F": 16, "G": 14}, freeze="A6")
    title(ws, "Abonnementen", "Checklist van alles wat maandelijks afschrijft. Zet 'Nog nodig?' op Nee en zie wat je bespaart.")
    header_row(ws, 5, 2, ["Dienst", "Prijs per maand", "Prijs per jaar", "Betaaldag", "Opzegtermijn", "Nog nodig?"])
    n = 20
    for i in range(n):
        r = 6 + i
        input_cell(ws, f"B{r}", align=al_left)
        input_cell(ws, f"C{r}", EUR)
        formula_cell(ws, f"D{r}", f"=C{r}*12", EUR)
        input_cell(ws, f"E{r}", "0", al_center)
        input_cell(ws, f"F{r}", align=al_center)
        input_cell(ws, f"G{r}", align=al_center)
    rt = 6 + n
    label_cell(ws, f"B{rt}", "Totaal", bold=True, fill=fill_sand)
    formula_cell(ws, f"C{rt}", f"=SUM(C6:C{rt - 1})", EUR, bold=True, fill=fill_sand)
    formula_cell(ws, f"D{rt}", f"=SUM(D6:D{rt - 1})", EUR, bold=True, fill=fill_sand)
    for col in "EFG":
        cell(ws, f"{col}{rt}", None, f_base, fill_sand, brd=border)
    label_cell(ws, f"B{rt + 1}", "Besparing als je alle 'Nee' opzegt", bold=True, fill=fill_sand)
    formula_cell(ws, f"C{rt + 1}", f'=SUMIF(G6:G{rt - 1},"Nee",C6:C{rt - 1})', EUR, bold=True, fill=fill_sand)
    formula_cell(ws, f"D{rt + 1}", f"=C{rt + 1}*12", EUR, bold=True, fill=fill_sand)
    for col in "EFG":
        cell(ws, f"{col}{rt + 1}", None, f_base, fill_sand, brd=border)
    ws.conditional_formatting.add(f"G6:G{rt - 1}", CellIsRule(operator="equal", formula=['"Nee"'], font=f_red))

    dv = DataValidation(type="list", formula1='"Ja,Nee"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G6:G{rt - 1}")
    cell(ws, f"B{rt + 3}", "Tip: zet de opzegtermijn (bv. '1 maand') erbij en plan de opzegging vóór de volgende betaaldag.", f_small)
    return ws


# ---------------------------------------------------------------- 52 weken
def build_52(wb):
    ws = wb.create_sheet("Spaarpotjes 52 weken")
    prep(ws, {"A": 3, "B": 8, "C": 16, "D": 10, "E": 3, "F": 16, "G": 10, "H": 3, "I": 16, "J": 10}, freeze="A9")
    title(ws, "Spaarpotjes 52 weken", "De 52-weken spaarchallenge: elke week iets meer opzij. Zet een ✓ als het gelukt is.")

    label_cell(ws, "B5", "Eigen variant – stap per week (€)", bold=True)
    ws.merge_cells("B5:D5")
    c = input_cell(ws, "F5", EUR)
    c.value = 2
    cell(ws, "G5", "Week 1 = 1× stap, week 52 = 52× stap. Vul een ander bedrag in voor een eigen challenge.", f_small)

    R_H = 8
    R0 = 9
    R_T = R0 + 52   # 61
    header_row(ws, R_H, 2, ["Week", "Klassiek (€1 → €52)", "Gedaan"])
    header_row(ws, R_H, 6, ["Omgekeerd (€52 → €1)", "Gedaan"])
    header_row(ws, R_H, 9, ["Eigen variant", "Gedaan"])
    for i in range(52):
        r = R0 + i
        cell(ws, f"B{r}", i + 1, f_bold, None, "0", al_center, border)
        formula_cell(ws, f"C{r}", f"=B{r}", EUR)
        input_cell(ws, f"D{r}", align=al_center)
        formula_cell(ws, f"F{r}", f"=53-B{r}", EUR)
        input_cell(ws, f"G{r}", align=al_center)
        formula_cell(ws, f"I{r}", f"=B{r}*$F$5", EUR)
        input_cell(ws, f"J{r}", align=al_center)
    label_cell(ws, f"B{R_T}", "Totaal", bold=True, fill=fill_sand)
    label_cell(ws, f"B{R_T + 1}", "Al gespaard", bold=True, fill=fill_sand)
    for amt, chk in (("C", "D"), ("F", "G"), ("I", "J")):
        formula_cell(ws, f"{amt}{R_T}", f"=SUM({amt}{R0}:{amt}{R_T - 1})", EUR, bold=True, fill=fill_sand)
        formula_cell(ws, f"{amt}{R_T + 1}", f'=SUMIF({chk}{R0}:{chk}{R_T - 1},"✓",{amt}{R0}:{amt}{R_T - 1})',
                     EUR, bold=True, fill=fill_sand)
        cell(ws, f"{chk}{R_T}", None, f_base, fill_sand, brd=border)
        cell(ws, f"{chk}{R_T + 1}", None, f_base, fill_sand, brd=border)

    dv = DataValidation(type="list", formula1='"✓"', allow_blank=True)
    ws.add_data_validation(dv)
    for chk in "DGJ":
        dv.add(f"{chk}{R0}:{chk}{R_T - 1}")
    ws.conditional_formatting.add(f"D{R0}:D{R_T - 1}", CellIsRule(operator="equal", formula=['"✓"'], fill=fill_green, font=f_head))
    ws.conditional_formatting.add(f"G{R0}:G{R_T - 1}", CellIsRule(operator="equal", formula=['"✓"'], fill=fill_green, font=f_head))
    ws.conditional_formatting.add(f"J{R0}:J{R_T - 1}", CellIsRule(operator="equal", formula=['"✓"'], fill=fill_green, font=f_head))
    cell(ws, f"B{R_T + 3}", "Klassiek en omgekeerd leveren allebei € 1.378,00 op na 52 weken. Omgekeerd is fijn als je in december liever weinig hoeft te missen.", f_small)
    return ws


# ---------------------------------------------------------------- main
def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    build_start(wb)
    build_instellingen(wb)
    build_jaar(wb)
    for m in MONTHS:
        build_month(wb, m)
    build_spaardoelen(wb)
    build_schulden(wb)
    build_abonnementen(wb)
    build_52(wb)
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = GREEN
    wb.worksheets[0].sheet_properties.tabColor = DARK
    wb.save(OUT_FILE)
    print("Opgeslagen:", OUT_FILE)


if __name__ == "__main__":
    main()
